#!/usr/bin/env python3
"""Rebuild deterministic gold JSON from the packaged synthetic inputs.

The examples intentionally publish their gold data. They teach task construction and
are not a hidden test set for leaderboard claims.
"""

from __future__ import annotations

import csv
import json
import re
import tarfile
import tempfile
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
TASK_ROOT = ROOT / "examples" / "harbor-office-tasks"


def _write(task: str, data: dict[str, Any]) -> None:
    path = TASK_ROOT / task / "tests" / "gold" / "gold_answer.json"
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _extract(task: str, destination: Path) -> Path:
    archive = TASK_ROOT / task / "environment" / "workspace.tar.gz"
    with tarfile.open(archive, "r:gz") as bundle:
        for member in bundle.getmembers():
            target = (destination / member.name).resolve()
            if destination.resolve() not in target.parents and target != destination.resolve():
                raise ValueError(f"unsafe archive member: {member.name}")
        bundle.extractall(destination, filter="data")
    return destination / "input"


def _parse_day(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for pattern in ("%Y/%m/%d", "%Y-%m-%d", "%Y%m%d"):
        try:
            return datetime.strptime(str(value), pattern).date()
        except ValueError:
            pass
    raise ValueError(value)


def _cleaning(input_dir: Path) -> dict[str, Any]:
    workbook = load_workbook(input_dir / "raw_fee_table.xlsx", data_only=True)
    aliases = {row[0]: (row[1], row[3]) for row in workbook["product_aliases"].iter_rows(min_row=2, values_only=True)}
    fee_records: list[dict[str, Any]] = []

    for row in workbook["sales_fee_export_raw"].iter_rows(min_row=2, values_only=True):
        product, share, fee_type, condition, raw_value, investor_group, _, note = row
        if "promotional waiver" in str(note or ""):
            continue
        if str(raw_value).endswith("%"):
            value, unit = float(str(raw_value).rstrip("%")), "percent"
        elif "1000" in str(raw_value):
            value, unit = 1000.0, "CNY_per_txn"
        elif "不收取" in str(raw_value):
            value, unit = 0.0, "percent"
        else:
            raise ValueError(f"unhandled sales fee: {raw_value!r}")
        product_id = aliases[product][0]
        fee_records.append({
            "id": product_id, "share": share, "section": "sales_fee", "type": fee_type,
            "investor_group": investor_group, "condition": condition, "value": value, "unit": unit,
        })

    for row in workbook["operation_fee_export_raw"].iter_rows(min_row=2, values_only=True):
        product, share, fee_type, raw_value, raw_unit, *_ = row
        if raw_value is None:
            continue
        value = float(raw_value) / 100 if raw_unit == "bp/year" else float(raw_value)
        fee_records.append({
            "id": aliases[product][0], "share": share, "section": "operation_fee", "type": fee_type,
            "investor_group": "全部投资群体", "condition": "annual_rate", "value": value, "unit": "percent_per_year",
        })

    for row in workbook["comprehensive_fee_export_raw"].iter_rows(min_row=2, values_only=True):
        product, share, raw_value, *_ = row
        if raw_value is None:
            continue
        fee_records.append({
            "id": aliases[product][0], "share": share, "section": "comprehensive_fee",
            "type": "comprehensive_fee_rate", "condition": "annualized_estimate",
            "investor_group": "全部投资群体", "value": float(raw_value), "unit": "percent_per_year",
        })

    return {
        "case_id": "office-fee-data-cleaning-L3-051",
        "case_type": "cleaning",
        "output_contract": {"path": "/workspace/output/result.xlsx", "type": "xlsx"},
        "required_sheets": ["产品对比", "费用明细", "风险等级", "异常数据", "处理说明"],
        "products": [
            {"id": "FP-1301", "share": "A", "name": "安桥稳收一年持有", "risk": "R2"},
            {"id": "FP-1301", "share": "C", "name": "安桥稳收一年持有", "risk": "R2"},
            {"id": "FP-1302", "share": "A", "name": "柏舟成长精选", "risk": "R4"},
            {"id": "FP-1303", "share": "A", "name": "澄岳均衡配置", "risk": "R3"},
        ],
        "fee_records": fee_records,
        "risk_records": [
            {"id": "FP-1301", "share": "A", "source_label": "稳健型客户可关注", "risk": "R2", "order": 2, "suitability": "稳健型及以上"},
            {"id": "FP-1301", "share": "C", "source_label": "中低风险", "risk": "R2", "order": 2, "suitability": "稳健型及以上"},
            {"id": "FP-1302", "share": "A", "source_label": "进取型", "risk": "R4", "order": 4, "suitability": "积极型及以上"},
            {"id": "FP-1303", "share": "A", "source_label": "平衡型", "risk": "R3", "order": 3, "suitability": "平衡型及以上"},
        ],
        "minimum_exception_rows": 6,
        "exception_token_groups": [
            ["管理费", "冲突"], ["综合费率", "缺失"], ["高波动成长型", "无法"], ["渠道优惠"],
        ],
        "effective_date": "2026-06-01",
    }


def _csv_rows(input_dir: Path) -> tuple[list[dict[str, str]], list[dict[str, str]]]:
    orders: list[dict[str, str]] = []
    receipts: list[dict[str, str]] = []
    for path in input_dir.glob("*.csv"):
        with path.open(encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.DictReader(handle))
        if rows and "ORDER_ID" in rows[0]:
            orders = rows
        elif rows and "RECEIPT_NO" in rows[0]:
            receipts = rows
    return orders, receipts


def _procurement(input_dir: Path) -> dict[str, Any]:
    orders, receipts = _csv_rows(input_dir)
    order_by_id = {row["ORDER_ID"]: row for row in orders}
    receipt_groups: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in receipts:
        receipt_groups[re.sub(r"-\d{2}$", "", row["RECEIPT_NO"])].append(row)

    matched, quantity, amount, single = [], [], [], []
    aggregate: dict[str, dict[str, float]] = {}
    for order_id, order in order_by_id.items():
        linked = receipt_groups.get(order_id, [])
        if not linked:
            single.append(order_id)
            continue
        receipt_qty = sum(float(row["RECV_QTY"]) for row in linked)
        receipt_amount = sum(float(row["RECV_AMOUNT"]) for row in linked)
        aggregate[order_id] = {
            "order_qty": float(order["ORDER_QTY"]), "receipt_qty": receipt_qty,
            "order_amount": float(order["TOTAL_AMOUNT"]), "receipt_amount": receipt_amount,
        }
        qty_diff = receipt_qty != float(order["ORDER_QTY"])
        amount_diff = abs(receipt_amount - float(order["TOTAL_AMOUNT"])) > 0.01
        if not qty_diff and not amount_diff:
            matched.append(order_id)
        if qty_diff:
            quantity.append(order_id)
        if amount_diff:
            amount.append(order_id)
    for base_id, linked in receipt_groups.items():
        if base_id not in order_by_id:
            single.extend(row["RECEIPT_NO"] for row in linked)

    sample_ids = {
        "数量差异": ["PO-2026-14490", "PO-2026-21715", "PO-2026-11619"],
        "金额差异": ["PO-2026-15944", "PO-2026-35540", "PO-2026-11619"],
    }
    sample_values = []
    for sheet, ids in sample_ids.items():
        for order_id in ids:
            sample_values.append({"sheet": sheet, "id": order_id, **aggregate[order_id]})

    return {
        "case_id": "office-procurement-reconcile-L3-052", "case_type": "reconciliation",
        "output_contract": {"path": "/workspace/output/result.xlsx", "type": "xlsx"},
        "required_sheets": ["对账汇总", "匹配记录", "数量差异", "金额差异", "单边记录", "数据质量", "处理说明"],
        "summary": {
            "订单总数": len(orders), "入库单总数": len(receipts), "完全匹配": len(matched),
            "数量差异": len(quantity), "金额差异": len(amount), "数量金额同时差异": len(set(quantity) & set(amount)),
            "未入库订单": sum(item in order_by_id for item in single), "无对应订单入库": sum(item not in order_by_id for item in single),
            "订单总金额": round(sum(float(row["TOTAL_AMOUNT"]) for row in orders), 2),
            "入库总金额": round(sum(float(row["RECV_AMOUNT"]) for row in receipts), 2),
        },
        "detail_ids": {
            "匹配记录": sorted(matched), "数量差异": sorted(quantity), "金额差异": sorted(amount), "单边记录": sorted(single),
        },
        "sample_values": sample_values,
    }


def _sheet_rows_by_kind(input_dir: Path) -> dict[str, tuple[list[str], list[tuple[Any, ...]]]]:
    result = {}
    for path in input_dir.glob("*.xlsx"):
        workbook = load_workbook(path, data_only=True, read_only=True)
        ws = workbook.active
        headers = [str(cell.value) for cell in ws[1]]
        kind = "settlement" if "COMMISSION" in headers else "payment"
        result[kind] = (headers, list(ws.iter_rows(min_row=2, values_only=True)))
        workbook.close()
    return result


def _anomaly(input_dir: Path) -> dict[str, Any]:
    sources = _sheet_rows_by_kind(input_dir)
    comparisons: dict[str, list[dict[str, Any]]] = {"结算对比": [], "回款对比": []}
    channels: set[str] = set()
    source_meta = {}
    for kind, (headers, raw_rows) in sources.items():
        rows = list(dict.fromkeys(raw_rows))
        source_meta[kind] = {"raw": len(raw_rows), "deduplicated": len(rows), "excluded_quarter": sum(row[headers.index("QUARTER")] in (None, "") for row in rows)}
        periods = ("25Q4", "26Q1") if kind == "settlement" else ("25Q1", "26Q1")
        grouped: dict[str, dict[str, list[tuple[Any, ...]]]] = defaultdict(lambda: defaultdict(list))
        for row in rows:
            quarter = row[headers.index("QUARTER")]
            if quarter in periods:
                grouped[str(row[headers.index("CHANNEL_NAME")])][str(quarter)].append(row)
        channels.update(grouped)
        for channel in sorted(grouped):
            def total(period: str, field: str) -> float:
                return sum(float(row[headers.index(field)] or 0) for row in grouped[channel][period])
            if kind == "settlement":
                metric_values = {
                    "GMV": (total(periods[0], "GMV"), total(periods[1], "GMV")),
                    "REVENUE": (total(periods[0], "REVENUE"), total(periods[1], "REVENUE")),
                    "COMMISSION_RATE": (total(periods[0], "COMMISSION") / total(periods[0], "GMV"), total(periods[1], "COMMISSION") / total(periods[1], "GMV")),
                    "ORDER_COUNT": (total(periods[0], "ORDER_COUNT"), total(periods[1], "ORDER_COUNT")),
                }
                target = "结算对比"
            else:
                metric_values = {
                    "GMV": (total(periods[0], "GMV"), total(periods[1], "GMV")),
                    "REVENUE": (total(periods[0], "REVENUE"), total(periods[1], "REVENUE")),
                    "ENTERPRISE_INCOME": (
                        total(periods[0], "REVENUE") - total(periods[0], "PLATFORM_FEE") - total(periods[0], "SERVICE_FEE"),
                        total(periods[1], "REVENUE") - total(periods[1], "PLATFORM_FEE") - total(periods[1], "SERVICE_FEE"),
                    ),
                    "HEADCOUNT": (
                        total(periods[0], "HEADCOUNT") / len(grouped[channel][periods[0]]),
                        total(periods[1], "HEADCOUNT") / len(grouped[channel][periods[1]]),
                    ),
                }
                target = "回款对比"
            for metric, (previous, current) in metric_values.items():
                difference = current - previous
                growth = difference / previous if previous else None
                comparisons[target].append({
                    "channel": channel, "metric": metric, "previous": previous, "current": current,
                    "difference": difference, "growth": growth,
                })

    anomalies = [
        item for rows in comparisons.values() for item in rows
        if item["growth"] is not None and abs(item["growth"]) > 0.20
    ]
    return {
        "case_id": "office-channel-anomaly-analysis-L4-054", "case_type": "anomaly_analysis",
        "output_contract": {"path": "/workspace/output/result.xlsx", "type": "xlsx"},
        "required_sheets": ["结算对比", "回款对比", "重点异常", "数据质量", "处理说明"],
        "channels": sorted(channels), "source_counts": source_meta,
        "comparisons": comparisons, "anomalies": anomalies,
    }


def _weekly(input_dir: Path) -> dict[str, Any]:
    path = next(input_dir.glob("*.xlsx"))
    workbook = load_workbook(path, data_only=True, read_only=True)
    ws = workbook.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    workbook.close()
    grouped: dict[int, list[Any]] = defaultdict(lambda: [0, 0.0, 0.0, None, None])
    for row in rows:
        day = _parse_day(row[0])
        week = day.isocalendar().week
        values = grouped[week]
        values[0] += 1
        values[1] += float(row[3])
        values[2] += float(row[4])
        values[3] = min(values[3] or day, day)
        values[4] = max(values[4] or day, day)
    summary = [{
        "week": week, "start": values[3].isoformat(), "end": values[4].isoformat(), "count": values[0],
        "plan": values[1], "resolved": values[2], "rate": values[2] / values[1],
    } for week, values in sorted(grouped.items())]
    return {
        "case_id": "office-ticket-weekly-dashboard-L3-053", "case_type": "weekly_dashboard",
        "output_contract": {"path": "/workspace/output/result.xlsx", "type": "xlsx"},
        "required_sheets": ["周汇总", "数据质量", "处理说明"], "total_input_rows": len(rows),
        "weekly_summary": summary,
        "totals": {"plan": sum(item["plan"] for item in summary), "resolved": sum(item["resolved"] for item in summary)},
    }


def _forecast(input_dir: Path) -> dict[str, Any]:
    payload = json.loads((input_dir / "monthly_demand.json").read_text(encoding="utf-8"))
    history = {item["month"]: {"demand": item["demand"], "promotion": item["promotion"], "stockout_days": item["stockout_days"]} for item in payload["records"]}
    validation_months = ["2025-10", "2025-11", "2025-12"]
    recent = [float(history[month]["demand"]) for month in sorted(history)[-12:] if history[month]["demand"] is not None]
    return {
        "case_id": "office-demand-forecast-L4-055", "case_type": "forecast",
        "output_contract": {"path": "/workspace/output/result.xlsx", "type": "xlsx"},
        "required_sheets": ["历史数据", "历史验证", "预测结果", "预测说明"],
        "history": history, "missing_month": "2024-06",
        "validation_actuals": {month: history[month]["demand"] for month in validation_months},
        "maximum_validation_mape": 0.25,
        "forecast_months": ["2026-01", "2026-02", "2026-03"],
        "forecast_bounds": [min(recent) * 0.75, max(recent) * 1.25],
        "interval_width_ratio": [0.05, 0.60], "algorithm_is_prescribed": False,
    }


def _governance() -> dict[str, Any]:
    return {
        "case_id": "office-governance-summary-L4-056", "case_type": "governance_docx",
        "output_contract": {"path": "/workspace/output/result.docx", "type": "docx"},
        "required_sections": ["管理摘要", "里程碑", "风险与阻塞", "决策清单", "数据冲突", "来源说明"],
        "current_items": [
            {"id": "M-101", "tokens": ["已完成", "2026-06-30", "v1.8"]},
            {"id": "M-205", "tokens": ["已完成初筛", "抽样复核", "2026-07-03"]},
            {"id": "M-310", "tokens": ["启动日已锁定", "2026-07-15"]},
            {"id": "M-206", "tokens": ["黄色关注", "8月第一周"]},
            {"id": "DEC-044", "tokens": ["待治理会确认", "2026-07-05"]},
            {"id": "M-407", "tokens": ["冻结", "2026-07-04"]},
            {"id": "RISK-118", "tokens": ["2026-07-12", "8月第一周"]},
        ],
        "source_markers": ["登记簿", "里程碑更新", ".xlsx"],
        "conflicts": [
            {"id": "M-101", "required_groups": [["待签署", "旧稿"], ["已完成"], ["采用", "口径"]]},
            {"id": "M-205", "required_groups": [["草案待抽样", "旧稿"], ["已完成初筛"], ["采用", "口径"]]},
            {"id": "M-206", "required_groups": [["橙色", "旧稿"], ["黄色关注"], ["采用", "口径"]]},
        ],
        "unsupported_owner_items": ["M-206", "RISK-118"],
        "false_claims": [
            {"id": "M-101", "claim": "待签署"}, {"id": "DEC-044", "claim": "已完成"},
            {"id": "M-206", "claim": "橙色"}, {"id": "RISK-118", "claim": "已解除"},
        ],
        "minimum_characters": 1200,
    }


def main() -> int:
    builders = {
        "office-fee-data-cleaning-L3-051": _cleaning,
        "office-procurement-reconcile-L3-052": _procurement,
        "office-ticket-weekly-dashboard-L3-053": _weekly,
        "office-channel-anomaly-analysis-L4-054": _anomaly,
        "office-demand-forecast-L4-055": _forecast,
    }
    with tempfile.TemporaryDirectory(prefix="agentic-office-evals-gold-") as temporary:
        root = Path(temporary)
        for task, builder in builders.items():
            _write(task, builder(_extract(task, root / task)))
    _write("office-governance-summary-L4-056", _governance())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
