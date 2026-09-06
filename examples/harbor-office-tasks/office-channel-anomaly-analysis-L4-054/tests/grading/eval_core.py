from __future__ import annotations

import json
import math
import os
import re
import shutil
import subprocess
import tempfile
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable, Sequence


# Gate checks carry no score. If any gate fails, the final reward is zero.
# Every non-gate rubric has a fixed denominator of 1.0 for its task type.
RUBRICS: dict[str, list[tuple[str, float, bool]]] = {
    "cleaning": [
        ("output_exists", 0.0, True), ("workbook_opens", 0.0, True),
        ("required_sheets", 0.0, True), ("required_headers", 0.10, False),
        ("product_records", 0.15, False), ("fee_records", 0.20, False),
        ("risk_mapping", 0.12, False), ("key_exceptions", 0.10, False),
        ("typed_data", 0.08, False), ("source_trace", 0.10, False),
        ("formula_integrity", 0.10, False), ("workbook_usability", 0.05, False),
    ],
    "reconciliation": [
        ("output_exists", 0.0, True), ("workbook_opens", 0.0, True),
        ("required_sheets", 0.0, True), ("required_headers", 0.08, False),
        ("summary_values", 0.15, False), ("detail_membership", 0.25, False),
        ("detail_unique", 0.08, False), ("sample_values", 0.14, False),
        ("numeric_types", 0.08, False), ("source_trace", 0.06, False),
        ("formula_integrity", 0.10, False), ("workbook_usability", 0.06, False),
    ],
    "weekly_dashboard": [
        ("output_exists", 0.0, True), ("workbook_opens", 0.0, True),
        ("required_sheets", 0.0, True), ("required_headers", 0.08, False),
        ("weekly_rows", 0.30, False), ("totals_reconcile", 0.10, False),
        ("data_quality_audit", 0.08, False), ("weekly_charts", 0.16, False),
        ("chart_bindings", 0.10, False), ("completion_rate_format", 0.05, False),
        ("helper_data_not_exposed", 0.03, False), ("formula_integrity", 0.07, False),
        ("workbook_usability", 0.03, False),
    ],
    "anomaly_analysis": [
        ("output_exists", 0.0, True), ("workbook_opens", 0.0, True),
        ("required_sheets", 0.0, True), ("required_headers", 0.08, False),
        ("channel_coverage", 0.12, False), ("comparison_aggregates", 0.22, False),
        ("dedup_audit", 0.08, False), ("excluded_rows_audit", 0.05, False),
        ("anomaly_membership", 0.17, False), ("anomaly_math", 0.12, False),
        ("anomaly_evidence_fields", 0.08, False), ("formula_integrity", 0.05, False),
        ("workbook_usability", 0.03, False),
    ],
    "forecast": [
        ("output_exists", 0.0, True), ("workbook_opens", 0.0, True),
        ("required_sheets", 0.0, True), ("history_records", 0.12, False),
        ("missing_value_handling", 0.07, False), ("context_fields", 0.05, False),
        ("validation_actuals", 0.15, False), ("validation_math", 0.12, False),
        ("validation_quality", 0.10, False), ("future_forecast", 0.08, False),
        ("data_driven_intervals", 0.08, False), ("forecast_chart", 0.08, False),
        ("forecast_explanation", 0.08, False), ("formula_integrity", 0.05, False),
        ("workbook_usability", 0.02, False),
    ],
    "governance_docx": [
        ("output_exists", 0.0, True), ("document_opens", 0.0, True),
        ("required_sections", 0.10, False), ("current_item_facts", 0.25, False),
        ("conflicts_explained", 0.12, False), ("scope_boundaries", 0.08, False),
        ("source_trace", 0.12, False), ("unsupported_owners_pending", 0.10, False),
        ("qualified_overall_status", 0.08, False), ("document_structure", 0.08, False),
        ("substantive_content", 0.04, False), ("false_claims_qualified", 0.03, False),
    ],
}


def _check(name: str, passed: bool, detail: str) -> dict[str, Any]:
    return {"name": name, "passed": bool(passed), "detail": detail}


def _finalize(case_type: str, observed: Sequence[dict[str, Any]]) -> dict[str, Any]:
    by_name = {item["name"]: dict(item) for item in observed}
    checks: list[dict[str, Any]] = []
    for name, weight, gate in RUBRICS[case_type]:
        item = by_name.get(name, _check(name, False, "not evaluated because a prerequisite failed"))
        item["weight"] = weight
        item["gate"] = gate
        checks.append(item)
    gate_failed = any(item["gate"] and not item["passed"] for item in checks)
    earned = sum(item["weight"] for item in checks if item["passed"])
    possible = sum(item["weight"] for item in checks)
    reward = 0.0 if gate_failed else earned / possible
    return {
        "reward": round(reward, 10),
        "earned_weight": round(earned, 10),
        "possible_weight": round(possible, 10),
        "gate_failed": gate_failed,
        "passed": sum(item["passed"] for item in checks),
        "total": len(checks),
        "checks": checks,
    }


def _norm(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"[\s：:，,（）()_\-/]+", "", str(value)).lower()


def _is_number(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(float(value))


def _near(value: Any, expected: float, tolerance: float = 1e-4) -> bool:
    return _is_number(value) and abs(float(value) - expected) <= tolerance * max(1.0, abs(expected))


def _row_text(row: Iterable[Any]) -> str:
    return " | ".join("" if value is None else str(value) for value in row)


def _rows(ws) -> list[list[Any]]:
    return [[cell.value for cell in row] for row in ws.iter_rows()]


def _workbook_text(wb) -> str:
    return "\n".join(_row_text(row) for ws in wb.worksheets for row in _rows(ws))


def _date_key(value: Any) -> str | None:
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    text = str(value or "").strip()
    for pattern in (r"(20\d{2})[-/.年](\d{1,2})[-/.月](\d{1,2})", r"(20\d{2})(\d{2})(\d{2})"):
        match = re.search(pattern, text)
        if match:
            year, month, day = (int(part) for part in match.groups())
            try:
                return date(year, month, day).isoformat()
            except ValueError:
                return None
    return None


def _month_key(value: Any) -> str | None:
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m")
    match = re.search(r"(20\d{2})[-/.年](\d{1,2})", str(value or ""))
    if not match:
        return None
    return f"{int(match.group(1)):04d}-{int(match.group(2)):02d}"


def _header_map(ws, aliases: dict[str, list[str]], max_rows: int = 12) -> tuple[int, dict[str, int]] | None:
    wanted = {key: [_norm(alias) for alias in values] for key, values in aliases.items()}
    for row_no, row in enumerate(ws.iter_rows(min_row=1, max_row=min(max_rows, ws.max_row)), 1):
        values = [_norm(cell.value) for cell in row]
        mapping: dict[str, int] = {}
        for key, names in wanted.items():
            found = next((index for index, value in enumerate(values) if value and any(value == name or name in value for name in names)), None)
            if found is not None:
                mapping[key] = found
        if len(mapping) == len(aliases):
            return row_no, mapping
    return None


def _data_rows(ws, header_row: int) -> list[list[Any]]:
    return [row for row in _rows(ws)[header_row:] if any(value not in (None, "") for value in row)]


def _label_number(ws, aliases: list[str]) -> float | None:
    wanted = [_norm(alias) for alias in aliases]
    for row in ws.iter_rows():
        for index, cell in enumerate(row):
            text = _norm(cell.value)
            if text and any(text == alias or alias in text for alias in wanted):
                for offset in range(1, 4):
                    if index + offset < len(row) and _is_number(row[index + offset].value):
                        return float(row[index + offset].value)
    return None


def _chart_objects(ws) -> list[Any]:
    charts = list(ws._charts)
    nested = [sub for chart in charts for sub in getattr(chart, "_charts", []) if sub is not chart]
    return charts + nested


def _chart_formula_text(ws) -> str:
    formulas: list[str] = []
    for chart in _chart_objects(ws):
        for series in getattr(chart, "ser", []):
            for attr in ("cat", "val", "xVal", "yVal"):
                ref = getattr(series, attr, None)
                if ref is not None:
                    formulas.append(str(ref))
    return "\n".join(formulas)


def _workbook_usability(wb, required: Sequence[str]) -> tuple[bool, str]:
    failures = []
    for name in required:
        ws = wb[name]
        if ws.max_row < 2 or ws.max_column < 2:
            failures.append(f"{name}:empty")
        if ws.sheet_state != "visible":
            failures.append(f"{name}:not-visible")
    return not failures, ", ".join(failures) or "required sheets are visible and populated"


def _recalculated_formula_errors(path: Path, formula_wb) -> tuple[list[str], str]:
    error_tokens = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!")
    errors: list[str] = []
    formula_count = 0
    for ws in formula_wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")):
                    formula_count += 1
                    formula = str(cell.value)
                    if any(token in formula for token in error_tokens) or re.search(r"/\s*[+-]?0+(?:\.0+)?(?![\d.])", formula):
                        errors.append(f"{ws.title}!{cell.coordinate}={formula}")
    if not formula_count:
        return errors, "no formulas present"

    executable = shutil.which("soffice") or shutil.which("libreoffice")
    if not executable:
        errors.append("recalculation_unavailable: soffice/libreoffice not found")
        return errors, f"formulas={formula_count}; recalculation unavailable"

    try:
        with tempfile.TemporaryDirectory(prefix="agentic-office-evals-recalc-") as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "source.xlsx"
            out_dir = tmp_path / "out"
            profile = tmp_path / "profile"
            out_dir.mkdir()
            profile.mkdir()
            shutil.copy2(path, source)
            command = [
                executable, "--headless", f"-env:UserInstallation={profile.as_uri()}",
                "--convert-to", "xlsx", "--outdir", str(out_dir), str(source),
            ]
            completed = subprocess.run(command, capture_output=True, text=True, timeout=90, check=False)
            recalculated = out_dir / "source.xlsx"
            if completed.returncode != 0 or not recalculated.is_file():
                errors.append(f"recalculation_failed: exit={completed.returncode}")
            else:
                from openpyxl import load_workbook

                value_wb = load_workbook(recalculated, data_only=True, read_only=True)
                for ws in value_wb.worksheets:
                    for row in ws.iter_rows():
                        for cell in row:
                            if isinstance(cell.value, str) and any(token in cell.value for token in error_tokens):
                                errors.append(f"{ws.title}!{cell.coordinate}={cell.value}")
    except Exception as exc:
        errors.append(f"recalculation_failed: {exc!r}")
    return errors, f"formulas={formula_count}; errors={errors[:10]}"


def _load_xlsx(output_path: str):
    from openpyxl import load_workbook

    path = Path(output_path)
    if not path.is_file():
        return path, None, None
    try:
        return path, load_workbook(path, data_only=False), None
    except Exception as exc:
        return path, None, exc


def _xlsx_start(output_path: str, gold: dict[str, Any]) -> tuple[Path, Any, list[dict[str, Any]]]:
    path, wb, error = _load_xlsx(output_path)
    checks = [
        _check("output_exists", path.is_file(), f"expected {path}"),
        _check("workbook_opens", wb is not None, "opened successfully" if wb else repr(error)),
    ]
    if wb is not None:
        required = gold["required_sheets"]
        present = all(name in wb.sheetnames for name in required)
        checks.append(_check("required_sheets", present, f"sheets={wb.sheetnames}"))
    return path, wb, checks


def _formula_check(path: Path, wb) -> dict[str, Any]:
    errors, detail = _recalculated_formula_errors(path, wb)
    return _check("formula_integrity", not errors, detail)


def _eval_cleaning(output_path: str, gold: dict[str, Any]) -> dict[str, Any]:
    case_type = "cleaning"
    path, wb, checks = _xlsx_start(output_path, gold)
    if wb is None or not all(name in wb.sheetnames for name in gold["required_sheets"]):
        return _finalize(case_type, checks)

    specs = {
        "产品对比": {
            "id": ["canonical_product_id", "产品编号"], "share": ["share_class", "份额类别"],
            "name": ["canonical_product_name", "产品名称"], "risk": ["standard_risk_level", "风险等级"],
            "source": ["source_trace", "来源追溯", "来源"],
        },
        "费用明细": {
            "id": ["canonical_product_id", "产品编号"], "share": ["share_class", "份额类别"],
            "section": ["fee_section", "费用板块"], "type": ["fee_type", "费用类型"],
            "investor": ["investor_group", "投资群体"], "condition": ["condition", "条件"], "value": ["fee_value", "费用数值"],
            "unit": ["fee_unit", "费用单位"], "date": ["effective_date", "生效日期"],
            "source": ["source_trace", "来源追溯", "来源"],
        },
        "风险等级": {
            "id": ["canonical_product_id", "产品编号"], "share": ["share_class", "份额类别"],
            "source_label": ["source_risk_label", "原始风险标签"], "risk": ["standard_risk_level", "标准风险等级"],
            "order": ["risk_order", "风险顺序"], "suitability": ["sales_suitability_bucket", "适配范围"],
            "source": ["mapping_source", "映射来源", "来源"],
        },
        "异常数据": {
            "type": ["exception_type", "异常类型"], "field": ["field_name", "字段名"],
            "action": ["resolution_action", "处理方式", "处理结果"], "source": ["source_trace", "来源追溯", "来源"],
        },
    }
    tables = {name: _header_map(wb[name], aliases) for name, aliases in specs.items()}
    checks.append(_check("required_headers", all(tables.values()), f"missing={[name for name, value in tables.items() if not value]}"))
    if not all(tables.values()):
        checks.append(_formula_check(path, wb))
        ok, detail = _workbook_usability(wb, gold["required_sheets"])
        checks.append(_check("workbook_usability", ok, detail))
        return _finalize(case_type, checks)

    product_header, product_cols = tables["产品对比"]
    product_rows = _data_rows(wb["产品对比"], product_header)
    product_actual = {
        (_norm(row[product_cols["id"]]), _norm(row[product_cols["share"]])):
        (_norm(row[product_cols["name"]]), _norm(row[product_cols["risk"]]))
        for row in product_rows if len(row) > max(product_cols.values()) and row[product_cols["id"]]
    }
    expected_products = {
        (_norm(item["id"]), _norm(item["share"])): (_norm(item["name"]), _norm(item["risk"]))
        for item in gold["products"]
    }
    checks.append(_check("product_records", product_actual == expected_products, f"actual_keys={sorted(product_actual)}"))

    fee_header, fee_cols = tables["费用明细"]
    fee_rows = [row for row in _data_rows(wb["费用明细"], fee_header) if len(row) > max(fee_cols.values()) and row[fee_cols["id"]]]
    def fee_key(row: list[Any]) -> tuple[str, ...]:
        return tuple(_norm(row[fee_cols[name]]) for name in ("id", "share", "section", "type", "investor", "condition", "unit"))
    actual_fees = {fee_key(row): row[fee_cols["value"]] for row in fee_rows}
    missing_fees, wrong_fees = [], []
    for item in gold["fee_records"]:
        key = tuple(_norm(item[name]) for name in ("id", "share", "section", "type", "investor_group", "condition", "unit"))
        if key not in actual_fees:
            missing_fees.append(key)
        elif not _near(actual_fees[key], float(item["value"]), 1e-6):
            wrong_fees.append((key, actual_fees[key], item["value"]))
    fee_ok = len(fee_rows) == len(gold["fee_records"]) and len(actual_fees) == len(fee_rows) and not missing_fees and not wrong_fees
    checks.append(_check("fee_records", fee_ok, f"rows={len(fee_rows)} missing={missing_fees[:5]} wrong={wrong_fees[:5]}"))

    risk_header, risk_cols = tables["风险等级"]
    risk_rows = [row for row in _data_rows(wb["风险等级"], risk_header) if len(row) > max(risk_cols.values()) and row[risk_cols["id"]]]
    missing_risks = []
    for item in gold["risk_records"]:
        found = any(
            _norm(row[risk_cols["id"]]) == _norm(item["id"])
            and _norm(row[risk_cols["share"]]) == _norm(item["share"])
            and _norm(row[risk_cols["source_label"]]) == _norm(item["source_label"])
            and _norm(row[risk_cols["risk"]]) == _norm(item["risk"])
            and _near(row[risk_cols["order"]], float(item["order"]), 1e-9)
            and _norm(item["suitability"]) in _norm(row[risk_cols["suitability"]])
            for row in risk_rows
        )
        if not found:
            missing_risks.append(item)
    checks.append(_check("risk_mapping", len(risk_rows) == len(gold["risk_records"]) and not missing_risks, f"missing={missing_risks}"))

    exception_header, exception_cols = tables["异常数据"]
    exception_rows = _data_rows(wb["异常数据"], exception_header)
    exception_text = _norm("\n".join(_row_text(row) for row in exception_rows))
    missing_groups = [group for group in gold["exception_token_groups"] if not all(_norm(token) in exception_text for token in group)]
    checks.append(_check("key_exceptions", len(exception_rows) >= gold["minimum_exception_rows"] and not missing_groups, f"rows={len(exception_rows)} missing={missing_groups}"))

    numeric_ok = all(_is_number(row[fee_cols["value"]]) for row in fee_rows)
    date_ok = all(_date_key(row[fee_cols["date"]]) == gold["effective_date"] for row in fee_rows)
    checks.append(_check("typed_data", numeric_ok and date_ok, f"numeric={numeric_ok} dates={date_ok}"))
    sources_ok = all(str(row[fee_cols["source"]] or "").strip() for row in fee_rows) and all(str(row[product_cols["source"]] or "").strip() for row in product_rows)
    trace_text = _workbook_text(wb).lower()
    sources_ok = sources_ok and "raw_fee_table.xlsx" in trace_text and ".pdf" in trace_text
    checks.append(_check("source_trace", sources_ok, "every product/fee row needs a source and both XLSX/PDF sources must be named"))
    checks.append(_formula_check(path, wb))
    ok, detail = _workbook_usability(wb, gold["required_sheets"])
    checks.append(_check("workbook_usability", ok, detail))
    return _finalize(case_type, checks)


def _eval_reconcile(output_path: str, gold: dict[str, Any]) -> dict[str, Any]:
    case_type = "reconciliation"
    path, wb, checks = _xlsx_start(output_path, gold)
    if wb is None or not all(name in wb.sheetnames for name in gold["required_sheets"]):
        return _finalize(case_type, checks)

    detail_sheets = list(gold["detail_ids"])
    specs = {
        name: {
            "id": ["订单号", "ORDER_ID", "关联订单号"],
            "order_qty": ["订单数量", "ORDER_QTY"], "receipt_qty": ["入库数量", "RECV_QTY"],
            "order_amount": ["订单金额", "TOTAL_AMOUNT"], "receipt_amount": ["入库金额", "RECV_AMOUNT"],
            "source": ["来源行号", "source_row", "原始行号"],
        } for name in detail_sheets if name != "单边记录"
    }
    specs["单边记录"] = {
        "id": ["原始单号", "订单号", "单号"], "category": ["单边类型", "分类", "记录类型"],
        "source": ["来源行号", "source_row", "原始行号"],
    }
    tables = {name: _header_map(wb[name], aliases) for name, aliases in specs.items()}
    checks.append(_check("required_headers", all(tables.values()), f"missing={[name for name, value in tables.items() if not value]}"))

    aliases = {
        "订单总数": ["订单总数", "采购订单数"], "入库单总数": ["入库单总数", "入库记录总数", "入库记录数"],
        "完全匹配": ["完全匹配", "匹配记录"], "数量差异": ["数量差异"], "金额差异": ["金额差异"],
        "数量金额同时差异": ["数量金额同时差异", "同时存在数量和金额差异", "重叠数量"],
        "未入库订单": ["未入库订单"], "无对应订单入库": ["无对应订单入库", "无订单入库", "未匹配入库"],
        "订单总金额": ["订单总金额"], "入库总金额": ["入库总金额"],
    }
    mismatches = []
    for label, expected in gold["summary"].items():
        value = _label_number(wb["对账汇总"], aliases[label])
        if value is None or not _near(value, float(expected), 1e-6):
            mismatches.append((label, value, expected))
    checks.append(_check("summary_values", not mismatches, f"mismatches={mismatches}"))

    actual_ids: dict[str, list[str]] = {}
    if all(tables.values()):
        for name, expected_ids in gold["detail_ids"].items():
            header, cols = tables[name]
            rows = _data_rows(wb[name], header)
            ids = [str(row[cols["id"]] or "").strip() for row in rows if len(row) > cols["id"] and str(row[cols["id"]] or "").strip().startswith("PO-")]
            if name == "单边记录":
                ids = [re.sub(r"-\d{2}$", "", item) if item not in expected_ids else item for item in ids]
            actual_ids[name] = ids
        membership_ok = all(set(actual_ids[name]) == set(expected) for name, expected in gold["detail_ids"].items())
        unique_ok = all(len(items) == len(set(items)) == len(gold["detail_ids"][name]) for name, items in actual_ids.items())
    else:
        membership_ok = unique_ok = False
    checks.append(_check("detail_membership", membership_ok, f"counts={{{', '.join(f'{k}:{len(v)}' for k, v in actual_ids.items())}}}"))
    checks.append(_check("detail_unique", unique_ok, "each expected business key must appear exactly once per detail sheet"))

    sample_failures = []
    if all(tables.values()):
        for sample in gold["sample_values"]:
            name = sample["sheet"]
            header, cols = tables[name]
            rows = _data_rows(wb[name], header)
            row = next((item for item in rows if sample["id"] in _row_text(item)), None)
            if row is None:
                sample_failures.append(f"{name}:{sample['id']}:missing")
                continue
            for field in ("order_qty", "receipt_qty", "order_amount", "receipt_amount"):
                if field in sample and (field not in cols or not _near(row[cols[field]], float(sample[field]), 1e-6)):
                    sample_failures.append(f"{name}:{sample['id']}:{field}")
    else:
        sample_failures.append("headers unavailable")
    checks.append(_check("sample_values", not sample_failures, f"failures={sample_failures}"))

    numeric_failures, source_failures = [], []
    if all(tables.values()):
        for name in detail_sheets:
            header, cols = tables[name]
            rows = _data_rows(wb[name], header)
            for row in rows:
                if len(row) <= max(cols.values()) or not str(row[cols["id"]] or "").startswith("PO-"):
                    continue
                for field in ("order_qty", "receipt_qty", "order_amount", "receipt_amount"):
                    if field in cols and row[cols[field]] not in (None, "") and not _is_number(row[cols[field]]):
                        numeric_failures.append(f"{name}:{row[cols['id']]}:{field}")
                if not str(row[cols["source"]] or "").strip():
                    source_failures.append(f"{name}:{row[cols['id']]}")
    checks.append(_check("numeric_types", not numeric_failures, f"failures={numeric_failures[:10]}"))
    checks.append(_check("source_trace", not source_failures, f"missing={source_failures[:10]}"))
    checks.append(_formula_check(path, wb))
    ok, detail = _workbook_usability(wb, gold["required_sheets"])
    checks.append(_check("workbook_usability", ok, detail))
    return _finalize(case_type, checks)


def _week_number(value: Any) -> int | None:
    if _is_number(value) and 1 <= int(value) <= 53:
        return int(value)
    match = re.search(r"(?:第)?\s*(\d{1,2})\s*周", str(value or ""))
    return int(match.group(1)) if match else None


def _eval_weekly(output_path: str, gold: dict[str, Any]) -> dict[str, Any]:
    case_type = "weekly_dashboard"
    path, wb, checks = _xlsx_start(output_path, gold)
    if wb is None or not all(name in wb.sheetnames for name in gold["required_sheets"]):
        return _finalize(case_type, checks)

    ws = wb["周汇总"]
    aliases = {
        "week": ["周次", "自然周"], "start": ["开始日期", "周开始"], "end": ["结束日期", "周结束"],
        "count": ["记录数", "工单数"], "plan": ["计划处理量", "计划量"],
        "resolved": ["实际解决量", "解决量", "实际量"], "rate": ["完成率"],
    }
    table = _header_map(ws, aliases)
    checks.append(_check("required_headers", table is not None, f"required={list(aliases)}"))
    actual: dict[int, list[Any]] = {}
    actual_row_numbers: dict[int, int] = {}
    if table:
        header, cols = table
        for row_no in range(header + 1, ws.max_row + 1):
            row = [cell.value for cell in ws[row_no]]
            if len(row) > max(cols.values()):
                week = _week_number(row[cols["week"]])
                if week is not None:
                    actual[week] = row
                    actual_row_numbers[week] = row_no
        failures = []
        for expected in gold["weekly_summary"]:
            row = actual.get(expected["week"])
            if row is None:
                failures.append(f"week {expected['week']}:missing")
                continue
            comparisons = [
                (_date_key(row[cols["start"]]) == expected["start"], "start"),
                (_date_key(row[cols["end"]]) == expected["end"], "end"),
                (_near(row[cols["count"]], expected["count"], 1e-9), "count"),
                (_near(row[cols["plan"]], expected["plan"], 1e-9), "plan"),
                (_near(row[cols["resolved"]], expected["resolved"], 1e-9), "resolved"),
                (_near(row[cols["rate"]], expected["rate"], 5e-6), "rate"),
            ]
            failures.extend(f"week {expected['week']}:{name}" for passed, name in comparisons if not passed)
        weekly_ok = set(actual) == {item["week"] for item in gold["weekly_summary"]} and not failures
    else:
        failures, weekly_ok = ["header missing"], False
    checks.append(_check("weekly_rows", weekly_ok, f"weeks={sorted(actual)} failures={failures}"))
    if table:
        plan_total = sum(float(actual[item["week"]][cols["plan"]]) for item in gold["weekly_summary"] if item["week"] in actual and _is_number(actual[item["week"]][cols["plan"]]))
        resolved_total = sum(float(actual[item["week"]][cols["resolved"]]) for item in gold["weekly_summary"] if item["week"] in actual and _is_number(actual[item["week"]][cols["resolved"]]))
        totals_ok = _near(plan_total, gold["totals"]["plan"], 1e-9) and _near(resolved_total, gold["totals"]["resolved"], 1e-9)
    else:
        totals_ok, plan_total, resolved_total = False, 0, 0
    checks.append(_check("totals_reconcile", totals_ok, f"plan={plan_total} resolved={resolved_total}"))

    audit_text = _norm("\n".join(_row_text(row) for name in ["数据质量", "处理说明"] for row in _rows(wb[name])))
    audit_ok = all(_norm(token) in audit_text for token in [str(gold["total_input_rows"]), "重复", "0", "周一", "周日"])
    checks.append(_check("data_quality_audit", audit_ok, "input row count, duplicate result, and Monday-Sunday rule must be recorded"))
    charts = _chart_objects(ws)
    series_count = sum(len(getattr(chart, "ser", [])) for chart in charts)
    titled = sum(bool(getattr(chart, "title", None)) for chart in charts)
    checks.append(_check("weekly_charts", len(charts) >= 2 and series_count >= 3 and titled >= 2, f"charts={len(charts)} series={series_count} titled={titled}"))
    bindings = _chart_formula_text(ws)
    binding_ok = bool(bindings) and "周汇总" in bindings and all(str(item["week"]) in _workbook_text(wb) for item in gold["weekly_summary"])
    checks.append(_check("chart_bindings", binding_ok, f"binding_chars={len(bindings)}"))
    rate_ok = bool(table) and all(
        _is_number(actual[item["week"]][cols["rate"]])
        and "%" in str(ws.cell(actual_row_numbers[item["week"]], cols["rate"] + 1).number_format)
        for item in gold["weekly_summary"] if item["week"] in actual
    )
    checks.append(_check("completion_rate_format", rate_ok and len(actual) == 4, "all four rates must be numeric percentage cells"))

    visible_helper = []
    for column_index in range(8, ws.max_column + 1):
        has_data = any(ws.cell(row, column_index).value not in (None, "") for row in range(1, ws.max_row + 1))
        letter = ws.cell(1, column_index).column_letter
        if has_data and not ws.column_dimensions[letter].hidden:
            visible_helper.append(letter)
    checks.append(_check("helper_data_not_exposed", not visible_helper, f"visible_helper_columns={visible_helper}"))
    checks.append(_formula_check(path, wb))
    ok, detail = _workbook_usability(wb, gold["required_sheets"])
    checks.append(_check("workbook_usability", ok, detail))
    return _finalize(case_type, checks)


def _comparison_row_candidates(ws, channel: str, metric: str) -> list[list[Any]]:
    metric_aliases = {
        "GMV": ["GMV"], "REVENUE": ["营收", "REVENUE"], "COMMISSION_RATE": ["佣金率", "COMMISSION_RATE"],
        "ORDER_COUNT": ["订单量", "订单数量", "ORDER_COUNT"], "ENTERPRISE_INCOME": ["企业端收益", "ENTERPRISE_INCOME"],
        "HEADCOUNT": ["人数", "HEADCOUNT"],
    }[metric]
    rows = [row for row in _rows(ws) if any(str(value or "").strip() == channel for value in row)]
    with_metric = [row for row in rows if any(_norm(alias) in _norm(_row_text(row)) for alias in metric_aliases)]
    return with_metric or rows


def _eval_anomaly(output_path: str, gold: dict[str, Any]) -> dict[str, Any]:
    case_type = "anomaly_analysis"
    path, wb, checks = _xlsx_start(output_path, gold)
    if wb is None or not all(name in wb.sheetnames for name in gold["required_sheets"]):
        return _finalize(case_type, checks)

    anomaly_aliases = {
        "channel": ["渠道商", "渠道名称"], "metric": ["数据项目", "指标"], "previous": ["前期值", "以前期间值"],
        "current": ["当前期值", "本期值"], "difference": ["相差值", "差额"], "growth": ["增减比例", "增长率", "变化率"],
        "rule": ["触发规则"], "evidence": ["数据依据", "依据"], "reason": ["可能原因"], "confirm": ["是否需要确认", "需要确认"],
    }
    comparison_aliases = {
        "channel": ["渠道商", "渠道名称"], "metric": ["数据项目", "指标"], "previous": ["前期值", "以前期间值"],
        "current": ["当前期值", "本期值"], "difference": ["相差值", "差额"], "growth": ["增减比例", "增长率", "变化率"],
    }
    tables = {
        "结算对比": _header_map(wb["结算对比"], comparison_aliases),
        "回款对比": _header_map(wb["回款对比"], comparison_aliases),
        "重点异常": _header_map(wb["重点异常"], anomaly_aliases),
    }
    checks.append(_check("required_headers", all(tables.values()), f"missing={[name for name, value in tables.items() if not value]}"))

    coverage_failures, aggregate_failures = [], []
    for sheet_name, expected_rows in gold["comparisons"].items():
        seen_channels = set()
        for expected in expected_rows:
            channel, metric = expected["channel"], expected["metric"]
            candidates = _comparison_row_candidates(wb[sheet_name], channel, metric)
            matched = any(
                all(any(_near(value, float(expected[field]), 2e-6) for value in row) for field in ("previous", "current", "difference", "growth"))
                for row in candidates
            )
            if not matched:
                aggregate_failures.append(f"{sheet_name}:{channel}:{metric}")
            else:
                seen_channels.add(channel)
        expected_channels = set(gold["channels"])
        if seen_channels != expected_channels:
            coverage_failures.append(f"{sheet_name}:found={len(seen_channels)}")
    checks.append(_check("channel_coverage", not coverage_failures, f"failures={coverage_failures}"))
    checks.append(_check("comparison_aggregates", not aggregate_failures, f"failures={aggregate_failures[:20]} total={len(aggregate_failures)}"))

    audit_text = _norm("\n".join(_row_text(row) for name in ["数据质量", "处理说明"] for row in _rows(wb[name])))
    dedup_ok = all(token in audit_text for token in ["1210", "1205", "1010", "1005", "5", "去重"])
    excluded_ok = all(token in audit_text for token in ["quarter", "空", "5", "不参与"])
    checks.append(_check("dedup_audit", dedup_ok, "must state raw/deduplicated row counts 1210/1205 and 1010/1005, plus five duplicates each"))
    checks.append(_check("excluded_rows_audit", excluded_ok, "must state that five blank QUARTER rows in each file were excluded"))

    anomaly_table = tables["重点异常"]
    actual_pairs: dict[tuple[str, str], list[Any]] = {}
    if anomaly_table:
        header, cols = anomaly_table
        for row in _data_rows(wb["重点异常"], header):
            if len(row) <= max(cols.values()):
                continue
            channel = str(row[cols["channel"]] or "").strip()
            metric_norm = _norm(row[cols["metric"]])
            canonical = next((name for name in ["GMV", "REVENUE", "COMMISSION_RATE", "ORDER_COUNT", "ENTERPRISE_INCOME", "HEADCOUNT"] if any(_norm(alias) in metric_norm for alias in {
                "GMV": ["GMV"], "REVENUE": ["营收", "REVENUE"], "COMMISSION_RATE": ["佣金率", "COMMISSION_RATE"],
                "ORDER_COUNT": ["订单量", "订单数量", "ORDER_COUNT"], "ENTERPRISE_INCOME": ["企业端收益", "ENTERPRISE_INCOME"], "HEADCOUNT": ["人数", "HEADCOUNT"],
            }[name])), None)
            if channel in gold["channels"] and canonical:
                actual_pairs[(channel, canonical)] = row
    expected_pairs = {(item["channel"], item["metric"]) for item in gold["anomalies"]}
    checks.append(_check("anomaly_membership", set(actual_pairs) == expected_pairs, f"found={len(actual_pairs)} expected={len(expected_pairs)} missing={list(expected_pairs-set(actual_pairs))[:10]} extra={list(set(actual_pairs)-expected_pairs)[:10]}"))
    math_failures, evidence_failures = [], []
    if anomaly_table:
        _, cols = anomaly_table
        expected_by_pair = {(item["channel"], item["metric"]): item for item in gold["anomalies"]}
        for pair, row in actual_pairs.items():
            expected = expected_by_pair.get(pair)
            if not expected:
                continue
            for field in ("previous", "current", "difference", "growth"):
                if not _near(row[cols[field]], float(expected[field]), 2e-6):
                    math_failures.append(f"{pair}:{field}")
            required_text = " ".join(str(row[cols[field]] or "") for field in ("rule", "evidence", "reason", "confirm"))
            if not str(row[cols["rule"]] or "").strip() or not str(row[cols["evidence"]] or "").strip() or "需要业务确认" not in required_text or "确认" not in str(row[cols["confirm"]] or ""):
                evidence_failures.append(str(pair))
    checks.append(_check("anomaly_math", not math_failures and len(actual_pairs) == len(expected_pairs), f"failures={math_failures[:20]}"))
    checks.append(_check("anomaly_evidence_fields", not evidence_failures and len(actual_pairs) == len(expected_pairs), f"failures={evidence_failures[:20]}"))
    checks.append(_formula_check(path, wb))
    ok, detail = _workbook_usability(wb, gold["required_sheets"])
    checks.append(_check("workbook_usability", ok, detail))
    return _finalize(case_type, checks)


def _find_month_row(ws, month: str) -> list[Any] | None:
    return next((row for row in _rows(ws) if any(_month_key(value) == month for value in row)), None)


def _eval_forecast(output_path: str, gold: dict[str, Any]) -> dict[str, Any]:
    case_type = "forecast"
    path, wb, checks = _xlsx_start(output_path, gold)
    if wb is None or not all(name in wb.sheetnames for name in gold["required_sheets"]):
        return _finalize(case_type, checks)

    history_aliases = {
        "month": ["月份", "month"], "original": ["原始需求", "原始值"], "processed": ["处理后需求", "处理后值"],
        "promotion": ["促销", "promotion"], "stockout": ["缺货天数", "stockout_days", "缺货"],
    }
    validation_aliases = {
        "month": ["月份", "month"], "prediction": ["预测值", "模拟预测值"], "actual": ["实际值", "actual"],
        "absolute_error": ["绝对误差"], "percentage_error": ["百分比误差", "误差率"],
    }
    forecast_aliases = {
        "month": ["月份", "month"], "forecast": ["预测值", "预测需求"], "lower": ["下限", "区间下限"], "upper": ["上限", "区间上限"],
    }
    history_table = _header_map(wb["历史数据"], history_aliases)
    validation_table = _header_map(wb["历史验证"], validation_aliases)
    forecast_table = _header_map(wb["预测结果"], forecast_aliases)

    history_failures = []
    if history_table:
        _, cols = history_table
        month_rows = {month: _find_month_row(wb["历史数据"], month) for month in gold["history"]}
        for month, expected in gold["history"].items():
            row = month_rows[month]
            if row is None or _month_key(row[cols["month"]]) != month:
                history_failures.append(f"{month}:missing")
                continue
            original = row[cols["original"]]
            if expected["demand"] is None:
                if original not in (None, ""):
                    history_failures.append(f"{month}:original-not-blank")
            elif not _near(original, expected["demand"], 1e-9):
                history_failures.append(f"{month}:original")
            if not _is_number(row[cols["processed"]]):
                history_failures.append(f"{month}:processed")
        exact_months = {_month_key(row[cols["month"]]) for row in _data_rows(wb["历史数据"], history_table[0]) if len(row) > cols["month"] and _month_key(row[cols["month"]])}
        history_ok = exact_months == set(gold["history"]) and not history_failures
        missing_row = month_rows[gold["missing_month"]]
        missing_ok = bool(missing_row and missing_row[cols["original"]] in (None, "") and _is_number(missing_row[cols["processed"]]))
        context_ok = all(
            row is not None and bool(row[cols["promotion"]]) == expected["promotion"] and _near(row[cols["stockout"]], expected["stockout_days"], 1e-9)
            for month, expected in gold["history"].items() for row in [month_rows[month]]
        )
    else:
        history_ok = missing_ok = context_ok = False
        history_failures.append("headers missing")
    checks.append(_check("history_records", history_ok, f"failures={history_failures[:20]}"))
    checks.append(_check("missing_value_handling", missing_ok, f"month={gold['missing_month']}"))
    checks.append(_check("context_fields", context_ok, "promotion and stockout fields must match all 36 source records"))

    validation_failures, errors = [], []
    if validation_table:
        _, cols = validation_table
        for month, actual_expected in gold["validation_actuals"].items():
            row = _find_month_row(wb["历史验证"], month)
            if row is None:
                validation_failures.append(f"{month}:missing")
                continue
            prediction, actual = row[cols["prediction"]], row[cols["actual"]]
            if not _near(actual, actual_expected, 1e-9) or not _is_number(prediction):
                validation_failures.append(f"{month}:values")
                continue
            absolute = abs(float(prediction) - float(actual))
            percentage = absolute / float(actual)
            if not _near(row[cols["absolute_error"]], absolute, 1e-6) or not _near(row[cols["percentage_error"]], percentage, 1e-6):
                validation_failures.append(f"{month}:math")
            errors.append(percentage)
    else:
        validation_failures.append("headers missing")
    checks.append(_check("validation_actuals", not any("missing" in item or "values" in item for item in validation_failures) and len(errors) == len(gold["validation_actuals"]), f"failures={validation_failures}"))
    checks.append(_check("validation_math", not any("math" in item for item in validation_failures) and len(errors) == len(gold["validation_actuals"]), f"failures={validation_failures}"))
    mape = sum(errors) / len(errors) if errors else math.inf
    checks.append(_check("validation_quality", mape <= gold["maximum_validation_mape"], f"MAPE={mape:.6f} limit={gold['maximum_validation_mape']}"))

    forecast_failures, interval_failures = [], []
    forecasts: list[float] = []
    if forecast_table:
        _, cols = forecast_table
        for month in gold["forecast_months"]:
            row = _find_month_row(wb["预测结果"], month)
            if row is None:
                forecast_failures.append(f"{month}:missing")
                continue
            values = [row[cols[name]] for name in ("forecast", "lower", "upper")]
            if not all(_is_number(value) for value in values):
                forecast_failures.append(f"{month}:non-numeric")
                continue
            forecast, lower, upper = (float(value) for value in values)
            forecasts.append(forecast)
            if not gold["forecast_bounds"][0] <= forecast <= gold["forecast_bounds"][1]:
                forecast_failures.append(f"{month}:outside-data-bounds")
            width_ratio = (upper - lower) / max(1.0, abs(forecast))
            if not (0 <= lower < forecast < upper and gold["interval_width_ratio"][0] <= width_ratio <= gold["interval_width_ratio"][1]):
                interval_failures.append(f"{month}:lower={lower},forecast={forecast},upper={upper}")
        result_months = {_month_key(row[cols["month"]]) for row in _data_rows(wb["预测结果"], forecast_table[0]) if len(row) > cols["month"] and _month_key(row[cols["month"]])}
        if not set(gold["forecast_months"]).issubset(result_months):
            forecast_failures.append("forecast months incomplete")
    else:
        forecast_failures.append("headers missing")
    checks.append(_check("future_forecast", not forecast_failures and len(forecasts) == 3, f"failures={forecast_failures}"))
    checks.append(_check("data_driven_intervals", not interval_failures and len(forecasts) == 3, f"failures={interval_failures}"))

    result_ws = wb["预测结果"]
    charts = _chart_objects(result_ws)
    series_count = sum(len(getattr(chart, "ser", [])) for chart in charts)
    bindings = _chart_formula_text(result_ws)
    blank_modes = [str(getattr(chart, "display_blanks", getattr(chart, "dispBlanksAs", ""))) for chart in charts]
    chart_ok = len(charts) >= 1 and series_count >= 3 and bool(bindings) and all("zero" not in value.lower() for value in blank_modes)
    checks.append(_check("forecast_chart", chart_ok, f"charts={len(charts)} series={series_count} bindings={len(bindings)} blanks={blank_modes}"))
    text = _norm(_workbook_text(wb))
    explanation_ok = all(token in text for token in ["验证", "选择", "限制"]) and ("不是承诺" in text or "不构成承诺" in text) and ("范围" in text or "区间" in text)
    checks.append(_check("forecast_explanation", explanation_ok, "must explain validation, choice, limitations, and that the interval is not a commitment"))
    checks.append(_formula_check(path, wb))
    ok, detail = _workbook_usability(wb, gold["required_sheets"])
    checks.append(_check("workbook_usability", ok, detail))
    return _finalize(case_type, checks)


def _date_token_present(text: str, token: str) -> bool:
    normalized = _norm(text)
    variants = [_norm(token)]
    match = re.fullmatch(r"(20\d{2})-(\d{2})-(\d{2})", token)
    if match:
        year, month, day = match.groups()
        variants.extend([_norm(f"{year}年{int(month)}月{int(day)}日"), _norm(f"{year}/{int(month)}/{int(day)}")])
    return any(variant in normalized for variant in variants)


def _eval_governance(output_path: str, gold: dict[str, Any]) -> dict[str, Any]:
    case_type = "governance_docx"
    from docx import Document

    path = Path(output_path)
    checks = [_check("output_exists", path.is_file(), f"expected {path}")]
    if not path.is_file():
        return _finalize(case_type, checks)
    try:
        doc = Document(path)
    except Exception as exc:
        checks.append(_check("document_opens", False, repr(exc)))
        return _finalize(case_type, checks)
    checks.append(_check("document_opens", True, "opened successfully"))
    paragraphs = [paragraph.text.strip() for paragraph in doc.paragraphs if paragraph.text.strip()]
    table_rows = [[cell.text.strip() for cell in row.cells] for table in doc.tables for row in table.rows]
    lines = paragraphs + [_row_text(row) for row in table_rows]
    text = "\n".join(lines)
    normalized = _norm(text)

    missing_sections = [section for section in gold["required_sections"] if _norm(section) not in normalized]
    checks.append(_check("required_sections", not missing_sections, f"missing={missing_sections}"))
    fact_failures, source_failures = [], []
    current_ids = [item["id"] for item in gold["current_items"]]
    for item in gold["current_items"]:
        item_lines = [line for line in lines if item["id"] in line and sum(other in line for other in current_ids) == 1]
        if not item_lines or not any(all(_date_token_present(line, token) for token in item["tokens"]) for line in item_lines):
            fact_failures.append(item["id"])
        if not any(any(marker in line for marker in gold["source_markers"]) for line in item_lines):
            source_failures.append(item["id"])
    checks.append(_check("current_item_facts", not fact_failures, f"failures={fact_failures}"))
    checks.append(_check("source_trace", not source_failures and "PPT" in text, f"item_sources_missing={source_failures}; PPT_named={'PPT' in text}"))

    conflict_failures = []
    for conflict in gold["conflicts"]:
        related = [line for line in lines if conflict["id"] in line]
        joined = " ".join(related)
        if not all(any(token in joined for token in group) for group in conflict["required_groups"]):
            conflict_failures.append(conflict["id"])
    checks.append(_check("conflicts_explained", not conflict_failures, f"failures={conflict_failures}"))
    boundary_terms = ["历史归档", "后续版本", "仅监控", "模板"]
    checks.append(_check("scope_boundaries", all(term in text for term in boundary_terms), f"required={boundary_terms}"))

    owner_failures = []
    for item_id in gold["unsupported_owner_items"]:
        related = [line for line in lines if item_id in line]
        if not related or not any("待确认" in line and "负责人" in line for line in related):
            owner_failures.append(item_id)
    checks.append(_check("unsupported_owners_pending", not owner_failures, f"failures={owner_failures}"))
    status_lines = [line for line in lines if "总体状态" in line]
    qualified = len(status_lines) == 1 and "黄色关注" in status_lines[0] and "综合判断" in status_lines[0]
    checks.append(_check("qualified_overall_status", qualified, f"status_lines={status_lines}"))
    heading_count = sum(bool(paragraph.text.strip()) and str(paragraph.style.name).lower().startswith("heading") for paragraph in doc.paragraphs)
    checks.append(_check("document_structure", heading_count >= 6 and len(doc.tables) >= 5, f"headings={heading_count} tables={len(doc.tables)}"))
    checks.append(_check("substantive_content", len(text) >= gold["minimum_characters"], f"characters={len(text)}"))
    false_failures = []
    for item in gold["false_claims"]:
        for line in [value for value in lines if item["id"] in value and item["claim"] in value]:
            if not any(marker in line for marker in ["旧稿", "PPT", "冲突", "不采用", "不再", "错误"]):
                false_failures.append(f"{item['id']}:{item['claim']}")
    checks.append(_check("false_claims_qualified", not false_failures, f"failures={false_failures}"))
    return _finalize(case_type, checks)


def evaluate(output_path: str, gold_path: str) -> dict[str, Any]:
    gold = json.loads(Path(gold_path).read_text(encoding="utf-8"))
    evaluators = {
        "cleaning": _eval_cleaning,
        "reconciliation": _eval_reconcile,
        "weekly_dashboard": _eval_weekly,
        "anomaly_analysis": _eval_anomaly,
        "forecast": _eval_forecast,
        "governance_docx": _eval_governance,
    }
    return evaluators[gold["case_type"]](output_path, gold)
