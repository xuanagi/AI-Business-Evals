from __future__ import annotations

import json
import shutil
from datetime import date
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference

from grading.eval_core import RUBRICS, evaluate


ROOT = Path(__file__).resolve().parents[1]
TASK_ROOT = ROOT / "examples" / "harbor-office-tasks"
TASKS = {
    "cleaning": "office-fee-data-cleaning-L3-051",
    "reconciliation": "office-procurement-reconcile-L3-052",
    "weekly_dashboard": "office-ticket-weekly-dashboard-L3-053",
    "anomaly_analysis": "office-channel-anomaly-analysis-L4-054",
    "forecast": "office-demand-forecast-L4-055",
    "governance_docx": "office-governance-summary-L4-056",
}


def gold_path(case_type: str) -> Path:
    return TASK_ROOT / TASKS[case_type] / "tests" / "gold" / "gold_answer.json"


def gold(case_type: str) -> dict:
    return json.loads(gold_path(case_type).read_text(encoding="utf-8"))


def save_workbook(path: Path, sheets: dict[str, list[list[object]]]) -> Workbook:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        ws = workbook.create_sheet(name)
        for row in rows:
            ws.append(row)
    workbook.save(path)
    return workbook


def build_cleaning(path: Path) -> None:
    expected = gold("cleaning")
    products = [["canonical_product_id", "share_class", "canonical_product_name", "standard_risk_level", "source_trace"]]
    products += [[item["id"], item["share"], item["name"], item["risk"], "summary.pdf"] for item in expected["products"]]
    fees = [["canonical_product_id", "share_class", "fee_section", "fee_type", "investor_group", "condition", "fee_value", "fee_unit", "effective_date", "source_trace"]]
    fees += [[item["id"], item["share"], item["section"], item["type"], item["investor_group"], item["condition"], item["value"], item["unit"], date(2026, 6, 1), "raw_fee_table.xlsx; summary.pdf"] for item in expected["fee_records"]]
    risks = [["canonical_product_id", "share_class", "source_risk_label", "standard_risk_level", "risk_order", "sales_suitability_bucket", "mapping_source"]]
    risks += [[item["id"], item["share"], item["source_label"], item["risk"], item["order"], item["suitability"], "risk_codebook.xlsx"] for item in expected["risk_records"]]
    exceptions = [
        ["exception_type", "field_name", "resolution_action", "source_trace"],
        ["resolved_effective_date_conflict", "管理费", "记录冲突并采用新日期", "raw_fee_table.xlsx"],
        ["missing_field", "综合费率", "缺失，不估算", "raw_fee_table.xlsx"],
        ["unmapped_risk", "高波动成长型", "无法映射", "ops_profile_notes.md"],
        ["fee_waiver_note", "渠道优惠", "只作提示", "raw_fee_table.xlsx"],
        ["missing_field", "管理费", "缺失", "raw_fee_table.xlsx"],
        ["scoped_out_field", "管理人", "不纳入", "standard_schema.json"],
    ]
    save_workbook(path, {"产品对比": products, "费用明细": fees, "风险等级": risks, "异常数据": exceptions, "处理说明": [["项目", "说明"], ["输入", "raw_fee_table.xlsx 与 summary.pdf"]]})


def build_reconciliation(path: Path) -> None:
    expected = gold("reconciliation")
    summary = [["指标", "数值"]] + [[label, value] for label, value in expected["summary"].items()]
    samples = {(item["sheet"], item["id"]): item for item in expected["sample_values"]}
    sheets: dict[str, list[list[object]]] = {"对账汇总": summary}
    headers = ["订单号", "订单数量", "入库数量", "订单金额", "入库金额", "来源行号"]
    for name in ["匹配记录", "数量差异", "金额差异"]:
        rows: list[list[object]] = [headers]
        for index, item_id in enumerate(expected["detail_ids"][name], 2):
            item = samples.get((name, item_id), {})
            rows.append([item_id, item.get("order_qty", 1), item.get("receipt_qty", 1), item.get("order_amount", 1), item.get("receipt_amount", 1), index])
        sheets[name] = rows
    sheets["单边记录"] = [["原始单号", "单边类型", "来源行号"]] + [[item_id, "未入库订单" if not item_id.endswith("-01") else "无对应订单入库", index] for index, item_id in enumerate(expected["detail_ids"]["单边记录"], 2)]
    sheets["数据质量"] = [["项目", "结果"], ["空白和重复", "已检查"]]
    sheets["处理说明"] = [["项目", "说明"], ["来源", "采购订单.csv；入库记录.csv"]]
    save_workbook(path, sheets)


def build_weekly(path: Path) -> None:
    expected = gold("weekly_dashboard")
    workbook = Workbook()
    ws = workbook.active
    ws.title = "周汇总"
    ws.append(["周次", "开始日期", "结束日期", "记录数", "计划处理量", "实际解决量", "完成率"])
    for item in expected["weekly_summary"]:
        ws.append([item["week"], date.fromisoformat(item["start"]), date.fromisoformat(item["end"]), item["count"], item["plan"], item["resolved"], item["rate"]])
        ws.cell(ws.max_row, 7).number_format = "0.0%"
    bar = BarChart()
    bar.title = "计划与实际"
    bar.add_data(Reference(ws, min_col=5, max_col=6, min_row=1, max_row=5), titles_from_data=True)
    bar.set_categories(Reference(ws, min_col=1, min_row=2, max_row=5))
    ws.add_chart(bar, "I2")
    line = LineChart()
    line.title = "完成率"
    line.add_data(Reference(ws, min_col=7, min_row=1, max_row=5), titles_from_data=True)
    line.set_categories(Reference(ws, min_col=1, min_row=2, max_row=5))
    ws.add_chart(line, "I18")
    quality = workbook.create_sheet("数据质量")
    quality.append(["项目", "结果"])
    quality.append(["输入行数", 560])
    quality.append(["重复", 0])
    notes = workbook.create_sheet("处理说明")
    notes.append(["规则", "说明"])
    notes.append(["自然周", "周一开始，周日结束"])
    workbook.save(path)


def build_anomaly(path: Path) -> None:
    expected = gold("anomaly_analysis")
    headers = ["渠道商", "数据项目", "前期值", "当前期值", "相差值", "增减比例"]
    sheets: dict[str, list[list[object]]] = {}
    for name, records in expected["comparisons"].items():
        sheets[name] = [headers] + [[item["channel"], item["metric"], item["previous"], item["current"], item["difference"], item["growth"]] for item in records]
    anomaly_headers = headers + ["触发规则", "数据依据", "可能原因", "是否需要确认"]
    sheets["重点异常"] = [anomaly_headers] + [[item["channel"], item["metric"], item["previous"], item["current"], item["difference"], item["growth"], "绝对变化超过20%", "输入明细去重汇总", "需要业务确认", "是，需要确认"] for item in expected["anomalies"]]
    sheets["数据质量"] = [["项目", "结果"], ["结算", "原始1210，重复5，去重后1205；QUARTER为空5条，不参与汇总"], ["回款", "原始1010，重复5，去重后1005；QUARTER为空5条，不参与汇总"]]
    sheets["处理说明"] = [["项目", "说明"], ["规则", "先去重，再排除空 QUARTER"]]
    save_workbook(path, sheets)


def build_forecast(path: Path) -> None:
    expected = gold("forecast")
    history = [["月份", "原始需求", "处理后需求", "促销", "缺货天数"]]
    for month, item in expected["history"].items():
        processed = item["demand"] if item["demand"] is not None else 1050
        history.append([month, item["demand"], processed, item["promotion"], item["stockout_days"]])
    validation = [["月份", "预测值", "实际值", "绝对误差", "百分比误差"]]
    for month, actual in expected["validation_actuals"].items():
        prediction = actual * 0.9
        validation.append([month, prediction, actual, abs(prediction - actual), abs(prediction - actual) / actual])
    forecast = [["月份", "预测值", "下限", "上限"], ["2026-01", 1050, 950, 1150], ["2026-02", 1100, 990, 1210], ["2026-03", 1200, 1080, 1320]]
    workbook = save_workbook(path, {"历史数据": history, "历史验证": validation, "预测结果": forecast, "预测说明": [["项目", "说明"], ["选择", "根据历史验证选择"], ["限制", "区间只是范围，不构成承诺，也不是承诺值"]]})
    result = workbook["预测结果"]
    chart = LineChart()
    chart.title = "历史验证与未来预测"
    chart.add_data(Reference(result, min_col=2, max_col=4, min_row=1, max_row=4), titles_from_data=True)
    chart.set_categories(Reference(result, min_col=1, min_row=2, max_row=4))
    result.add_chart(chart, "F2")
    workbook.save(path)


def build_governance(path: Path) -> None:
    expected = gold("governance_docx")
    document = Document()
    document.add_heading("管理摘要", level=1)
    document.add_paragraph("总体状态：黄色关注（综合判断）。资料口径为 v1.8，日期为 2026-07-01，背景参考 PPT。")
    document.add_heading("里程碑", level=1)
    table = document.add_table(rows=1, cols=4)
    for cell, value in zip(table.rows[0].cells, ["事项编号", "状态与日期", "负责人", "来源"]):
        cell.text = value
    owners = {"M-206": "负责人：待确认", "RISK-118": "负责人：待确认"}
    for item in expected["current_items"]:
        cells = table.add_row().cells
        cells[0].text = item["id"]
        cells[1].text = "；".join(item["tokens"])
        cells[2].text = owners.get(item["id"], "负责人：登记簿所列团队")
        cells[3].text = "来源：program_milestone_update_register.xlsx / 里程碑更新"
    document.add_heading("风险与阻塞", level=1)
    risk = document.add_table(rows=2, cols=2)
    risk.cell(0, 0).text, risk.cell(0, 1).text = "事项", "说明"
    risk.cell(1, 0).text, risk.cell(1, 1).text = "M-206 / RISK-118", "黄色关注，负责人：待确认；来源：登记簿里程碑更新"
    document.add_heading("决策清单", level=1)
    decision = document.add_table(rows=2, cols=2)
    decision.cell(0, 0).text, decision.cell(0, 1).text = "事项", "决策"
    decision.cell(1, 0).text, decision.cell(1, 1).text = "DEC-044", "待治理会确认；来源：登记簿里程碑更新"
    document.add_heading("数据冲突", level=1)
    conflict = document.add_table(rows=1, cols=2)
    conflict.cell(0, 0).text, conflict.cell(0, 1).text = "事项", "冲突与采用口径"
    conflict_rows = {
        "M-101": "旧稿写待签署；当前已完成；本报告采用登记簿当前口径。",
        "M-205": "旧稿写草案待抽样；当前已完成初筛并进入抽样复核；本报告采用登记簿当前口径。",
        "M-206": "旧稿写橙色风险；当前为黄色关注；本报告采用登记簿当前口径。",
    }
    for item_id, text in conflict_rows.items():
        cells = conflict.add_row().cells
        cells[0].text, cells[1].text = item_id, text
    document.add_heading("范围边界", level=1)
    boundary = document.add_table(rows=2, cols=2)
    boundary.cell(0, 0).text, boundary.cell(0, 1).text = "类型", "处理"
    boundary.cell(1, 0).text, boundary.cell(1, 1).text = "历史归档、后续版本、仅监控、模板", "不替换当前治理口径"
    document.add_heading("来源说明", level=1)
    sources = document.add_table(rows=2, cols=2)
    sources.cell(0, 0).text, sources.cell(0, 1).text = "来源", "用途"
    sources.cell(1, 0).text, sources.cell(1, 1).text = "登记簿里程碑更新与 PPT", "当前事实与旧稿冲突核对"
    document.add_paragraph("本报告逐项保留编号、当前状态、日期与来源。" + "所有结论均以登记簿当前记录为准，并将旧稿冲突单列供复核。" * 60)
    document.save(path)


BUILDERS = {
    "cleaning": build_cleaning,
    "reconciliation": build_reconciliation,
    "weekly_dashboard": build_weekly,
    "anomaly_analysis": build_anomaly,
    "forecast": build_forecast,
    "governance_docx": build_governance,
}


@pytest.mark.parametrize("case_type", sorted(TASKS))
def test_missing_output_is_a_zero_score_with_fixed_rubric(tmp_path: Path, case_type: str) -> None:
    result = evaluate(str(tmp_path / "missing.file"), str(gold_path(case_type)))
    assert result["reward"] == 0.0
    assert result["gate_failed"] is True
    assert len(result["checks"]) == len(RUBRICS[case_type])


@pytest.mark.parametrize("case_type", sorted(BUILDERS))
def test_oracle_fixture_passes_all_checks(tmp_path: Path, case_type: str) -> None:
    suffix = ".docx" if case_type == "governance_docx" else ".xlsx"
    output = tmp_path / f"result{suffix}"
    BUILDERS[case_type](output)
    result = evaluate(str(output), str(gold_path(case_type)))
    failed = [item for item in result["checks"] if not item["passed"]]
    assert result["reward"] == 1.0, failed


def test_magic_numbers_without_keyed_week_rows_fail(tmp_path: Path) -> None:
    output = tmp_path / "result.xlsx"
    expected = gold("weekly_dashboard")
    magic = [["说明", "数值"]] + [["magic", value] for item in expected["weekly_summary"] for value in (item["plan"], item["resolved"], item["rate"])]
    save_workbook(output, {"周汇总": magic, "数据质量": [["项目", "结果"], ["重复", 0]], "处理说明": [["项目", "说明"], ["规则", "周一至周日"]]})
    result = evaluate(str(output), str(gold_path("weekly_dashboard")))
    assert next(item for item in result["checks"] if item["name"] == "weekly_rows")["passed"] is False


def test_invalid_formula_is_detected_without_cached_values(tmp_path: Path) -> None:
    output = tmp_path / "result.xlsx"
    build_weekly(output)
    from openpyxl import load_workbook

    workbook = load_workbook(output)
    workbook["处理说明"]["C2"] = "=1/0"
    workbook.save(output)
    result = evaluate(str(output), str(gold_path("weekly_dashboard")))
    assert next(item for item in result["checks"] if item["name"] == "formula_integrity")["passed"] is False


@pytest.mark.skipif(not (shutil.which("soffice") or shutil.which("libreoffice")), reason="LibreOffice is only required in the verifier image")
def test_libreoffice_recalculation_detects_runtime_formula_error(tmp_path: Path) -> None:
    output = tmp_path / "result.xlsx"
    build_weekly(output)
    from openpyxl import load_workbook

    workbook = load_workbook(output)
    workbook["处理说明"]["C2"] = "=SQRT(-1)"
    workbook.save(output)
    result = evaluate(str(output), str(gold_path("weekly_dashboard")))
    formula_check = next(item for item in result["checks"] if item["name"] == "formula_integrity")
    assert formula_check["passed"] is False
    assert "#NUM!" in formula_check["detail"]


def test_duplicate_procurement_ids_cannot_replace_missing_ids(tmp_path: Path) -> None:
    output = tmp_path / "result.xlsx"
    build_reconciliation(output)
    from openpyxl import load_workbook

    workbook = load_workbook(output)
    ws = workbook["匹配记录"]
    ws.cell(ws.max_row, 1).value = ws.cell(2, 1).value
    workbook.save(output)
    result = evaluate(str(output), str(gold_path("reconciliation")))
    assert next(item for item in result["checks"] if item["name"] == "detail_membership")["passed"] is False
    assert next(item for item in result["checks"] if item["name"] == "detail_unique")["passed"] is False


def test_forecast_with_arbitrary_large_values_fails_data_bounds(tmp_path: Path) -> None:
    output = tmp_path / "result.xlsx"
    build_forecast(output)
    from openpyxl import load_workbook

    workbook = load_workbook(output)
    for row in range(2, 5):
        workbook["预测结果"].cell(row, 2).value = 4999
        workbook["预测结果"].cell(row, 3).value = 4000
        workbook["预测结果"].cell(row, 4).value = 5000
    workbook.save(output)
    result = evaluate(str(output), str(gold_path("forecast")))
    assert next(item for item in result["checks"] if item["name"] == "future_forecast")["passed"] is False


def test_keyword_stuffed_governance_document_fails_item_association(tmp_path: Path) -> None:
    output = tmp_path / "result.docx"
    expected = gold("governance_docx")
    document = Document()
    document.add_paragraph(" ".join(expected["required_sections"]))
    document.add_paragraph(" ".join(token for item in expected["current_items"] for token in [item["id"], *item["tokens"]]))
    document.add_paragraph("历史归档 后续版本 仅监控 模板 登记簿 PPT 总体状态 黄色关注 综合判断")
    document.save(output)
    result = evaluate(str(output), str(gold_path("governance_docx")))
    assert next(item for item in result["checks"] if item["name"] == "current_item_facts")["passed"] is False


def test_partial_anomaly_spot_checks_do_not_pass_full_oracle(tmp_path: Path) -> None:
    output = tmp_path / "result.xlsx"
    expected = gold("anomaly_analysis")
    first = expected["comparisons"]["结算对比"][0]
    row = [first["channel"], first["metric"], first["previous"], first["current"], first["difference"], first["growth"]]
    headers = ["渠道商", "数据项目", "前期值", "当前期值", "相差值", "增减比例"]
    save_workbook(output, {"结算对比": [headers, row], "回款对比": [headers, row], "重点异常": [[*headers, "触发规则", "数据依据", "可能原因", "是否需要确认"], [*row, "20%", "输入", "需要业务确认", "是"]], "数据质量": [["项目", "结果"], ["去重", "5"]], "处理说明": [["项目", "说明"], ["规则", "去重"]]})
    result = evaluate(str(output), str(gold_path("anomaly_analysis")))
    assert next(item for item in result["checks"] if item["name"] == "comparison_aggregates")["passed"] is False


def test_cleaning_spot_rows_do_not_pass_complete_fee_oracle(tmp_path: Path) -> None:
    output = tmp_path / "result.xlsx"
    build_cleaning(output)
    from openpyxl import load_workbook

    workbook = load_workbook(output)
    ws = workbook["费用明细"]
    ws.delete_rows(7, ws.max_row - 6)
    workbook.save(output)
    result = evaluate(str(output), str(gold_path("cleaning")))
    assert next(item for item in result["checks"] if item["name"] == "fee_records")["passed"] is False
